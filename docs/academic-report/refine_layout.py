"""Refine an existing report without regenerating its frozen source inventory.
Supply a rendered PDF to refresh cached navigation page numbers. Run again
against the new render until page mappings stabilize, then inspect all pages.
"""
from pathlib import Path
import argparse,json,re,shutil,unicodedata
import pdfplumber
from docx import Document
from docx.shared import Pt,Twips,RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
parser=argparse.ArgumentParser()
parser.add_argument('--pdf',required=True)
args=parser.parse_args()
path=BASE/'CoVie_BaoCaoDoAn.docx'
backup=ROOT/'.local/report-qa-original/CoVie_BaoCaoDoAn.docx'
backup.parent.mkdir(parents=True,exist_ok=True)
if not backup.exists():shutil.copy2(path,backup)
doc=Document(path)

def norm(s):return re.sub(r'\s+','',unicodedata.normalize('NFC',s)).casefold()
with pdfplumber.open(args.pdf) as pdf:
    pages=[norm(p.extract_text() or '') for p in pdf.pages]

def page_for(text):
    term=norm(text)
    found=[i+1 for i,s in enumerate(pages) if term in s]
    if not found:raise ValueError('Heading/caption not found in PDF: '+text)
    return found[-1]

for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    if st.type not in [1,2]:continue
    rf=st.element.find('.//'+qn('w:rFonts'))
    if rf is not None:
        for a in list(rf.attrib):
            if 'Theme' in a:del rf.attrib[a]
        if st.name!='Code Block':
            for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
    if st.name.startswith('Heading') or st.name in ['Title','Subtitle','Header']:
        st.font.color.rgb=RGBColor(0,0,0)
        col=st.element.find('.//'+qn('w:color'))
        if col is not None:
            for a in list(col.attrib):
                if 'theme' in a.lower():del col.attrib[a]
for p in doc.paragraphs:
    for border in list(p._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
    if p.text=='BÁO CÁO ĐỒ ÁN':p.style=doc.styles['Title']
    if p.text.startswith('Mục lục điện tử bên dưới'):
        p.text='Mục lục có liên kết và số trang theo bản dàn trang đã kiểm tra. Khi chỉnh sửa nội dung, cần cập nhật lại các trường mục lục trước khi in.'
        p.style=doc.styles['Source Note']
    if p.style.name=='Caption' and p.text.startswith('Hình '):p.paragraph_format.keep_with_next=True

# More balanced tabular columns and vertically centered cells.
for tb in doc.tables:
    heads=[c.text for c in tb.rows[0].cells]
    widths=[int(x.get(qn('w:w'))) for x in tb._tbl.tblGrid]
    if len(heads)==3 and heads[0]=='Mã':widths=[1100,3900,4071]
    if len(heads)==2 and heads[0]=='Migration':widths=[7471,1600]
    for grid,w in zip(tb._tbl.tblGrid,widths):grid.set(qn('w:w'),str(w))
    borders=tb._tbl.tblPr.find(qn('w:tblBorders'))
    if borders is not None:
        for el in borders:el.set(qn('w:color'),'D9D9D9')
    for ri,row in enumerate(tb.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Twips(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if heads[ci] in ['Mã','Method','Số CREATE TABLE']:
                for p in cell.paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.CENTER

body=doc.element.body
entries={}
for p in doc.paragraphs:
    bookmarks=p._p.findall(qn('w:bookmarkStart'))
    if p.style.name.startswith('Heading') and bookmarks:
        entries[bookmarks[0].get(qn('w:name'))]={'text':p.text,'page':page_for(p.text)}

caps={}
for i,p in enumerate([p for p in doc.paragraphs if p.style.name=='Caption'],2000):
    name='caption_'+str(i)
    if not p._p.findall(qn('w:bookmarkStart')):
        b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(i));b.set(qn('w:name'),name)
        e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(i));p._p.insert(0,b);p._p.append(e)
    else:name=p._p.find(qn('w:bookmarkStart')).get(qn('w:name'))
    entries[name]={'text':p.text,'page':page_for(p.text)}
    caps[p.text]=name

def update_nav(p,name,title,pg):
    # Retain paragraph properties (indent/style), replace displayed runs only.
    for el in list(p):
        if el.tag!=qn('w:pPr'):p.remove(el)
    pr=p.find(qn('w:pPr'))
    if pr is None:pr=OxmlElement('w:pPr');p.insert(0,pr)
    tabs=pr.find(qn('w:tabs'))
    if tabs is not None:pr.remove(tabs)
    tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'right');tab.set(qn('w:leader'),'dot');tab.set(qn('w:pos'),'9071');tabs.append(tab);pr.append(tabs)
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),name)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=title;r.append(t);link.append(r);p.append(link)
    r=OxmlElement('w:r');r.append(OxmlElement('w:tab'));p.append(r)
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),' PAGEREF '+name+' \\h ')
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=str(pg);r.append(t);f.append(r);p.append(f)

for p in body.findall(qn('w:p')):
    style=p.find('w:pPr/w:pStyle',{'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
    if style is None or style.get(qn('w:val'))!='TOCEntry':continue
    link=p.find(qn('w:hyperlink'))
    if link is not None:
        name=link.get(qn('w:anchor'))
        if name in entries:
            e=entries[name];update_nav(p,name,e['text'],e['page'])
    else:
        title=''.join(t.text or '' for t in p.iter(qn('w:t')))
        if title in caps:
            name=caps[title];update_nav(p,name,title,entries[name]['page'])

from editorial_revision import apply_editorial
apply_editorial(doc)
doc.save(path)
(BASE/'navigation-pages.json').write_text(json.dumps({'source_pdf':str(Path(args.pdf).resolve()),'entries':entries},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'refined':str(path),'navigation_entries':len(entries),'source_pdf_pages':len(pages)},ensure_ascii=False))
