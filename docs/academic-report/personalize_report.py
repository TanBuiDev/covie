"""Apply user-supplied cover details to the existing paginated report."""
from pathlib import Path
import json,hashlib,shutil
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
logo=BASE/'assets/hutech-logo.png'
shutil.copy2(ROOT/'.local/hutech-01.png',logo)
details={'school':'TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP.HCM','student':'Bùi Lê Hoàng Tấn','student_id':'2280602879','class':'22DTHH2','cover_date':'07/2026','location':'Hồ Chí Minh','advisor':'omitted at user request','faculty':'Khoa Công nghệ thông tin','major':'Công nghệ phần mềm','logo_source':'https://file1.hutech.edu.vn/file/editor/homepage/stories/hinh34/logo%20CMYK-01.png','logo_source_page':'https://www.hutech.edu.vn/homepage/tin-t-c/ho-t-d-ng-sinh-vien/1744-y-nghia-logo-d-i-h-c-cong-ngh-tp-hcm-hutech','logo_sha256':hashlib.sha256(logo.read_bytes()).hexdigest()}
(BASE/'cover-metadata.json').write_text(json.dumps(details,ensure_ascii=False,indent=2),encoding='utf-8')
out=BASE/'CoVie_BaoCaoDoAn.docx'
backup=ROOT/'.local/report-before-hutech.docx'
if not backup.exists():shutil.copy2(out,backup)
doc=Document(out)
repl={'[TÊN TRƯỜNG]':details['school'],'[HỌ VÀ TÊN]':details['student'],'[MSSV]':details['student_id'],'[TÊN LỚP]':details['class'],'[ĐỊA ĐIỂM], [THÁNG/NĂM BẢO VỆ]':'Hồ Chí Minh, 07/2026'}
old_front='Bản thảo có đầy đủ nội dung kỹ thuật theo nguồn hiện có của dự án, phục vụ sinh viên rà soát và hoàn thiện trước khi nộp. Thông tin trường, khoa, sinh viên, giảng viên, tên đề tài chính thức và mẫu trình bày của khoa chưa được cung cấp; các vị trí tương ứng được đánh dấu trong ngoặc vuông. Không có chữ ký hoặc xác nhận nào được giả lập.'
new_front='Bản thảo phục vụ Bùi Lê Hoàng Tấn, MSSV 2280602879, lớp 22DTHH2, Trường Đại học Công nghệ TP.HCM (HUTECH) rà soát trước khi nộp. Bìa ghi Hồ Chí Minh, 07/2026 theo thông tin cung cấp; không hiển thị giảng viên hướng dẫn theo yêu cầu. Khoa Công nghệ thông tin, chuyên ngành Công nghệ phần mềm. Mẫu định dạng của khoa chưa được cung cấp. Lời cam đoan và lời cảm ơn cần sinh viên tự hoàn thiện.'
old_end='Thông tin chưa được cung cấp gồm trường, khoa, chuyên ngành, sinh viên, mã số sinh viên, lớp, giảng viên hướng dẫn, địa điểm và thời điểm bảo vệ. Các trường này ở bìa được để trong ngoặc vuông. Lời cam đoan cần chính người nộp kiểm tra và xác nhận; bản thảo không giả lập chữ ký hoặc xác nhận của nhà trường.'
new_end='Thông tin bìa do người dùng cung cấp: Bùi Lê Hoàng Tấn, MSSV 2280602879, lớp 22DTHH2; Hồ Chí Minh, 07/2026. Trường: Đại học Công nghệ TP.HCM (HUTECH). Không hiển thị giảng viên hướng dẫn theo yêu cầu. Khoa Công nghệ thông tin, chuyên ngành Công nghệ phần mềm. Logo lấy từ website chính thức HUTECH. Mốc trên bìa không thay đổi ngày tổng hợp và ngày kiểm chứng nguồn ghi trong nội dung. Sinh viên tự hoàn thiện lời cam đoan và lời cảm ơn.'
for p in list(doc.paragraphs):
    if p.text.startswith('Giảng viên hướng dẫn:'):
        p._p.getparent().remove(p._p);continue
    if p.text==old_front:p.text=new_front
    elif p.text==old_end:p.text=new_end
    else:
        for r in p.runs:
            for a,b in repl.items():
                if a in r.text:r.text=r.text.replace(a,b)
    if p.text==details['school']:
        for r in p.runs:r.font.size=Pt(14)
    if p.text=='Chuyên ngành: Công nghệ phần mềm':p.paragraph_format.space_after=Pt(12)
title=next(p for p in doc.paragraphs if p.text=='BÁO CÁO ĐỒ ÁN')
if not any(s._inline.docPr.get('descr','').startswith('Logo HUTECH') for s in doc.inline_shapes):
    p=title.insert_paragraph_before(style='Cover Meta');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(22)
    shape=p.add_run().add_picture(str(logo),width=Inches(3.1))
    shape._inline.docPr.set('descr','Logo HUTECH từ website chính thức của Trường Đại học Công nghệ TP.HCM')
doc.core_properties.author=details['student']
doc.save(out)
# Keep editable sources and future builds consistent without rebuilding the frozen inventory.
for path in [BASE/'build_report.py',BASE/'CoVie_BaoCaoDoAn.md',BASE/'chapters/08-phu-luc-tu-dong.md']:
    s=path.read_text(encoding='utf-8').replace(old_front,new_front).replace(old_end,new_end)
    if path.name=='build_report.py':
        for a,b in repl.items():s=s.replace(a,b)
        s=s.replace(",'Giảng viên hướng dẫn: [HỌC HÀM/HỌC VỊ, HỌ TÊN]'",'')
        s=s.replace("center('Chuyên ngành: Công nghệ phần mềm',13,after=44)","center('Chuyên ngành: Công nghệ phần mềm',13,after=12)\np=center('',after=22)\nshape=p.add_run().add_picture(str(BASE/'assets/hutech-logo.png'),width=Inches(3.1))\nshape._inline.docPr.set('descr','Logo HUTECH từ website chính thức của trường')")
    path.write_text(s,encoding='utf-8')
print('Personalized report with official unmodified logo')
