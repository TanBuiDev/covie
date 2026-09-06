"""Add nine UML use-case views without regenerating application inventories."""
from pathlib import Path
import json, hashlib, shutil
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from product_expansion import bookmark, navigation
from terminology import quote

BASE=Path(__file__).resolve().parent
ASSETS=BASE/'assets/usecases'
DATA=[
 dict(key='buyer',title='Người mua và hành trình giao dịch',actor='Người mua',cases=['Tìm kiếm và xem mặt hàng','Quản lý địa chỉ nhận hàng','Quản lý giỏ và tạo đơn','Theo dõi và phản hồi đơn','Đăng nhu cầu Cần mua','Đánh giá và xem thông báo'],
 before='Người mua sử dụng CoVie để tìm nguồn cung, tạo giao dịch và theo dõi kết quả. Sơ đồ tổng hợp các mục tiêu chính của vai trò này; chức năng khám phá còn được mở cho khách chưa đăng nhập, trong khi dữ liệu địa chỉ, giỏ, đơn và nhu cầu gắn với tài khoản.',
 after='UC-02 mô tả tìm kiếm và lựa chọn mặt hàng; UC-03 mô tả xem trước và tạo đơn. UC-04 và UC-05 làm rõ tương tác sau đặt hàng và hành trình Cần mua. Quyền đánh giá phụ thuộc giao dịch đủ điều kiện; thông báo và dữ liệu cá nhân chỉ được đọc trong phạm vi của người sử dụng. [N4, N8]'),
 dict(key='seller',title='Người bán và tổ chức hoạt động',actor='Thành viên\nngười bán',cases=['Quản lý hồ sơ và điểm bán','Quản lý mặt hàng và chào bán','Cập nhật giá và tồn kho','Cấu hình điều kiện giao nhận','Tiếp nhận và xử lý đơn','Phản hồi nhu cầu Cần mua'],
 before='Thành viên người bán tổ chức nguồn cung theo đơn vị bán hàng và điểm phục vụ. Sơ đồ tách việc chuẩn bị mặt hàng, điều kiện phục vụ và xử lý giao dịch thành các mục tiêu riêng. Một mặt hàng có thể được chào bán tại nhiều điểm với giá và tình trạng khác nhau.',
 after='Mỗi chức năng yêu cầu quyền thành viên và phạm vi điểm bán phù hợp. Sơ đồ thể hiện phạm vi yêu cầu của vai trò, còn mức độ kết nối giao diện và hợp đồng được đánh giá tại Chương 6. Đặc tả UC-04 bao quát xử lý đơn; UC-05 bao quát phản hồi nhu cầu. Người bán không được truy cập đơn hoặc thay đổi nguồn cung của đơn vị khác. [N4, N6, N9]'),
 dict(key='admin',title='Quản trị viên và kiểm soát hệ thống',actor='Quản trị viên',cases=['Tra cứu tài khoản người dùng','Hạn chế hoặc khôi phục tài khoản','Xem xét báo cáo vi phạm','Kiểm duyệt nội dung','Thay đổi cấu hình hệ thống','Tra cứu đơn và nhật ký'],
 before='Quản trị viên có các mục tiêu kiểm soát tài khoản, nội dung và chính sách vận hành. Tra cứu đơn hỗ trợ đối chiếu thông tin giao dịch; nhật ký giúp xác định người thực hiện, thời điểm và đối tượng của thay đổi.',
 after='Thao tác quản trị chịu ràng buộc quyền, trạng thái và phiên bản dữ liệu. Việc tra cứu đơn không trao quyền tùy ý chuyển trạng thái thay người mua hoặc người bán. UC-06 trình bày luồng kiểm duyệt và cập nhật cấu hình. Những trang chưa có hợp đồng quản trị chuyên biệt vẫn là giới hạn triển khai, dù thuộc phạm vi thiết kế giao diện. [N4, N9]'),
 dict(key='uc01',title='UC-01 Đăng nhập và quản lý phiên',actor='Người dùng',secondary='“Google”',right=[0],cases=['Đăng ký qua “Google”','Thiết lập mật khẩu','Đăng nhập bằng mật khẩu','Xem các phiên đăng nhập','Thu hồi phiên truy cập','Đăng xuất'],
 before='Sơ đồ phân rã UC-01 thành các mục tiêu xác lập và kiểm soát quyền truy cập. “Google” là tác nhân hệ thống bên ngoài tham gia xác minh danh tính khi đăng ký; kho phiên và cơ sở dữ liệu thuộc bên trong CoVie, không được biểu diễn như tác nhân.',
 after='Đăng ký lần đầu gồm xác minh danh tính và thiết lập mật khẩu theo luồng hiện có; sơ đồ không bổ sung đăng ký độc lập bằng “email”. Sau đăng nhập, người dùng có thể xem và thu hồi phiên theo quyền. Các trường hợp sai mật khẩu, phiên hết hiệu lực và tài khoản bị hạn chế được xử lý theo đặc tả tại mục 3.5. Đường liên kết trên hình biểu thị sự tham gia, không biểu thị thứ tự các bước. [N11]'),
 dict(key='uc02',title='UC-02 Tìm và chọn mặt hàng',actor='Khách',secondary='Người mua',left=[0,1,2,3,4],right=[0,1,2,3,4,5],cases=['Duyệt danh mục công khai','Tìm kiếm mặt hàng','Lọc và sắp xếp kết quả','Xem chi tiết mặt hàng','Xem hồ sơ người bán','Thêm lựa chọn vào giỏ'],
 before='Khách và người mua cùng tham gia khám phá nguồn cung công khai. Người mua đã xác thực có thêm mục tiêu lưu lựa chọn vào giỏ của mình. Sơ đồ giúp phân biệt chức năng có thể truy cập trước đăng nhập với chức năng phát sinh dữ liệu giao dịch cá nhân.',
 after='Chi tiết mặt hàng phân biệt hàng hóa, dịch vụ và đồ thanh lý. Không tìm thấy kết quả, lỗi tải dữ liệu và mặt hàng không còn khả dụng là các tình huống thay thế trong cùng chức năng. Kết quả tìm kiếm không bảo đảm giá hoặc lượng hàng còn nguyên khi đặt mua; dữ liệu được kiểm tra tiếp ở UC-03. Các tiêu chí bộ lọc và lựa chọn được trình bày tại mục 3.6. [N4, N7]'),
 dict(key='uc03',title='UC-03 Xem trước và tạo đơn',actor='Người mua',cases=['Chọn nguồn mua hợp lệ','Chọn hình thức giao nhận','Chọn địa chỉ nhận hàng','Chọn phương thức thanh toán','Xem trước nội dung và tiền','Xác nhận tạo đơn'],
 before='UC-03 tập trung vào quyết định đặt hàng của người mua. Nguồn mua có thể đến từ giỏ hoặc phản hồi Cần mua còn hiệu lực. Chọn địa chỉ chỉ áp dụng khi hình thức giao nhận yêu cầu; phương thức thanh toán thể hiện lựa chọn theo hợp đồng hiện có.',
 after='Địa chỉ và nguồn mua phải thuộc phạm vi được sử dụng. Máy chủ tính lại dữ liệu khi xem trước và kiểm tra lại trước tạo đơn. Nếu giá, giỏ, địa chỉ hoặc lượng hàng thay đổi, người mua cần xem lại bản tính; hệ thống không tự chấp nhận thay. Kết quả thành công là đơn chờ xác nhận, chưa khẳng định đã thanh toán. Trình tự và ngoại lệ được đặc tả tại mục 3.7; sơ đồ này chỉ biểu diễn phạm vi tham gia. [N7]'),
 dict(key='uc04',title='UC-04 Xử lý và hoàn tất đơn',actor='Thành viên\nngười bán',secondary='Người mua',left=[0,1,2,4,5],right=[0,3,4,5],cases=['Xem đơn và hành động hợp lệ','Xác nhận hoặc từ chối đơn','Cập nhật giao hàng hoặc thực hiện','Xác nhận nhận hàng hoặc hoàn tất','Yêu cầu và phản hồi hủy đơn','Báo hoặc xử lý sự cố đơn'],
 before='UC-04 có sự tham gia của hai phía giao dịch. Người bán tiếp nhận và cập nhật quá trình thực hiện; người mua theo dõi, xác nhận kết quả hoặc phản hồi. Các mục tiêu hủy và xử lý sự cố được đặt trong cùng phạm vi vì quyền của mỗi phía thay đổi theo tình trạng đơn.',
 after='Liên kết tới nhóm chức năng không có nghĩa hai tác nhân được thực hiện mọi lệnh như nhau. Chẳng hạn, yêu cầu hủy và phản hồi yêu cầu hủy là các hành động khác nhau trong cùng mục tiêu; chỉ hành động được máy chủ cho phép mới hợp lệ. Giao hàng, đến nhận và dịch vụ có nhánh thực hiện riêng. Mục 3.8 và sơ đồ trạng thái tại Chương 4 mô tả các điều kiện chuyển tiếp. [N4, N7]'),
 dict(key='uc05',title='UC-05 Đăng nhu cầu và chọn phản hồi',actor='Người mua',secondary='Thành viên\nngười bán',left=[0,1,2,5],right=[3,4],cases=['Đăng nhu cầu Cần mua','Xem nhu cầu và các phản hồi','Hủy nhu cầu khi được phép','Xem nhu cầu phù hợp điểm bán','Gửi hoặc rút phản hồi','Chọn phản hồi để tạo đơn'],
 before='Người mua công bố nhu cầu khi chưa tìm được nguồn cung phù hợp. Thành viên người bán tiếp cận nhu cầu trong phạm vi được phép và phản hồi bằng chào bán hợp lệ. Người mua lựa chọn phản hồi còn có thể sử dụng để tiếp tục sang UC-03.',
 after='Hủy nhu cầu, rút phản hồi và chuyển sang đặt hàng đều phụ thuộc thời hạn, quyền và trạng thái hiện tại. Tạo đơn chưa tự động chứng minh nhu cầu đã được đáp ứng; mốc ghi nhận phải theo quy tắc liên kết nhu cầu với đơn. Nguồn dữ liệu chọn ứng viên phản hồi trên giao diện người bán còn hạn chế, được nêu tại mục 3.9 và Chương 6. [N8, N9]'),
 dict(key='uc06',title='UC-06 Kiểm duyệt và thay đổi cấu hình',actor='Quản trị viên',cases=['Xem hồ sơ báo cáo','Xem thông tin đối tượng liên quan','Ẩn hoặc khôi phục nội dung','Đóng hồ sơ khi được phép','Cập nhật cấu hình có phiên bản','Đọc nhật ký thay đổi'],
 before='UC-06 tách mục tiêu xem xét hồ sơ, thực hiện quyết định kiểm duyệt và quản lý chính sách. Quản trị viên cần đọc thông tin hiện tại trước khi thao tác; nhật ký phục vụ đối chiếu sau thay đổi. Sơ đồ không gán tác vụ kiểm duyệt cho người mua hoặc người bán.',
 after='Các thao tác có tác động cần lý do, xác nhận và quyền tương ứng. Cập nhật cấu hình yêu cầu giá trị hợp lệ cùng phiên bản đang có hiệu lực; khi xảy ra xung đột, quản trị viên phải đọc lại trước khi tiếp tục. Nhật ký thuộc phạm vi tra cứu, không phải nơi sửa lịch sử. Quyền quản trị và giới hạn hợp đồng được trình bày tại mục 3.10. [N4, N9]'),
]

def draw_diagram(data):
    # Plain UML associations; no arrows suggesting sequence or invented dependencies.
    scale=2;w,h=1000,800
    im=Image.new('RGB',(w*scale,h*scale),'white');dr=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/times.ttf',26*scale)
    bold=ImageFont.truetype('C:/Windows/Fonts/timesbd.ttf',27*scale)
    def line(points,width=2):dr.line([(int(x*scale),int(y*scale)) for x,y in points],fill='black',width=width*scale)
    def txt(x,y,s,f=font):
        dr.multiline_text((x*scale,y*scale),s,font=f,fill='black',anchor='mm',align='center',spacing=4*scale)
    def ellipse(box):dr.ellipse(tuple(int(x*scale) for x in box),outline='black',width=2*scale,fill='white')
    def actor(x,name):
        ellipse((x-19,329,x+19,367));line([(x,367),(x,428)]);line([(x-34,390),(x+34,390)]);line([(x,428),(x-30,477)]);line([(x,428),(x+30,477)]);txt(x,510,name)
    dr.rectangle((240*scale,20*scale,760*scale,775*scale),outline='black',width=2*scale)
    txt(500,48,'CoVie',bold)
    ys=[125+117*i for i in range(6)]
    for i in data.get('left',range(6)):line([(125,390),(285,ys[i])])
    if data.get('secondary'):
        for i in data.get('right',[]):line([(875,390),(715,ys[i])])
    for y,label in zip(ys,data['cases']):
        # Wrap Vietnamese at word boundaries to keep every oval readable.
        words=label.split();lines=['']
        for word in words:
            candidate=(lines[-1]+' '+word).strip()
            if dr.textlength(candidate,font=font)>390*scale and lines[-1]:lines.append(word)
            else:lines[-1]=candidate
        ellipse((285,y-43,715,y+43));txt(500,y,'\n'.join(lines))
    actor(95,data['actor'])
    if data.get('secondary'):actor(905,data['secondary'])
    ASSETS.mkdir(parents=True,exist_ok=True);path=ASSETS/(data['key']+'.png');im.save(path)
    return path

def apply_usecases(doc):
    end=next(p._p for p in doc.paragraphs if p.style.name=='Heading 1' and p.text.startswith('CHƯƠNG 4.'))
    old=next((p._p for p in doc.paragraphs if p.text.startswith('3.12. Sơ đồ') and p.style.name=='Heading 2'),None)
    if old is not None:
        while old is not end:nxt=old.getnext();old.getparent().remove(old);old=nxt
    def para(text,style='Normal',page=False):
        p=doc.add_paragraph(quote(text),style);p.paragraph_format.page_break_before=page;end.addprevious(p._p);return p
    heading=para('3.12. Sơ đồ “use case” theo vai trò và nghiệp vụ','Heading 2',True);bookmark(heading,'usecase_diagrams',7100)
    intro='Các sơ đồ sau bổ sung góc nhìn chức năng cho đặc tả UC-01 đến UC-06. Tác nhân nằm ngoài khung CoVie; hình bầu dục biểu diễn mục tiêu sử dụng và đường liền biểu diễn sự tham gia. Các liên kết không biểu diễn trình tự xử lý hay xác nhận mức độ hoàn tất triển khai.'
    para(intro);md=['','## '+heading.text,'',intro];manifest=[]
    for i,data in enumerate(DATA,1):
        title=f'3.12.{i}. '+data['title'];p=para(title,'Heading 3',i>1);bookmark(p,'usecase_section_'+str(i),7100+i)
        para(data['before']);path=draw_diagram(data)
        p=para('');p.alignment=1;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1
        picture=p.add_run().add_picture(str(path),width=Inches(5.8));picture._inline.docPr.set('descr','Sơ đồ “use case”: '+data['title'])
        caption=f'Hình 3.{i+1}. Sơ đồ “use case” '+data['title'];p=para(caption,'Caption');p.paragraph_format.keep_with_next=True;bookmark(p,'usecase_figure_'+str(i),7200+i)
        source='Nguồn: mô hình hóa từ yêu cầu và đặc tả chức năng CoVie.';para(source,'Source Note');para(data['after'])
        md+=['','### '+quote(title),'',quote(data['before']),'',f'[[USECASE_FIGURE:{data["key"]}]]','',quote(caption),'',source,'',quote(data['after'])]
        manifest.append({'key':data['key'],'caption':quote(caption),'path':str(path.relative_to(BASE)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    navigation(doc)
    cp=BASE/'chapters/03-phan-tich-yeu-cau.md';original=cp.read_text(encoding='utf-8').split('\n## 3.12. Sơ đồ')[0].rstrip();cp.write_text(original+'\n'+'\n'.join(md)+'\n',encoding='utf-8')
    (BASE/'usecase-diagrams.json').write_text(json.dumps({'views':DATA,'assets':manifest},ensure_ascii=False,indent=2),encoding='utf-8')

if __name__=='__main__':
    f=BASE/'CoVie_BaoCaoDoAn.docx';backup=BASE.parents[1]/'.local/report-before-usecases.docx'
    if not backup.exists():shutil.copy2(f,backup)
    doc=Document(f);apply_usecases(doc);doc.save(f)
    body='\n\n'.join(p.read_text(encoding='utf-8') for p in sorted((BASE/'chapters').glob('*.md')))
    (BASE/'CoVie_BaoCaoDoAn.md').write_text('# CoVie — Báo cáo đồ án\n\n'+body,encoding='utf-8')
    print('Added nine use-case diagrams and corresponding explanations')
