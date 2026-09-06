"""Consistent double quotation of retained foreign technical terminology."""
from pathlib import Path
import re, json, shutil
from docx import Document

BASE = Path(__file__).resolve().parent
# Longest phrases win; identifiers/commands are never translated.
TERMS = '''Spring Boot|React Hook Form|TanStack Query|App Router|Server Components|Client Components|OAuth2 Resource Server|OAuth2 Client|Data JPA|Data Redis|Maven Wrapper|Docker Compose|Server-Sent Events|Cache API|service worker|refresh token|access token|Access JWT|Cookie HttpOnly|Bearer header|Public Buyer P0|Public Buyer|Buyer Mobile|Seller Desktop|Desktop Seller|Seller Center|Order Explorer|Shops & Locations|Product Detail|Checkout Success|Shop Profile|Address List|Service Detail|Resale Detail|test case|use case|location variant|rating stats|reputation stats|item snapshot|contact snapshot|policy snapshot|shipping evaluator|shipping fee|base fee|user ID|HTTP method|REST state|invalidation best-effort|fixture-gated skip|clean verify|production build|Exit code|SELECT FOR UPDATE|CREATE TABLE|NOT NULL|EXPLAIN (ANALYZE, BUFFERS)|S3-compatible|Project Status|Next.js|Node.js|PostgreSQL|PostGIS|OpenAPI|TypeScript|JavaScript|React|Java|Spring|Redis|MinIO|Flyway|Figma|ChatGPT|Codex|Google|Zod|Maven|Vitest|Playwright|Docker|Dockerfile|Hikari|Actuator|Validation|Security|Web|API|HTTP|SQL|PWA|SSE|ERD|JWT|CSRF|OAuth2|OAuth|DTO|URL|ID|UUID|RSA|CRUD|REST|UI|JDK|LTS|DB|COD|PK|FK|UNIQUE|NULL|ALTER|TABLE|GET|POST|PATCH|DELETE|PUT|ETag|If-Match|Idempotency-Key|EventSource|HttpOnly|pnpm|marketplace|checkout changed|checkout|preview|commit|backend|frontend|local|production|staging|migration|snapshot|catalog|client|outbox|token|fixture|media|merchant|merchants|Buyer|Seller|Admin|buyer|seller|admin|auth|users|orders|needs|reviews|categories|checkouts|location-areas|need-requests|reports|settings|location|locations|delivery|shipping|inventory|cart|carts|reputation|notifications|moderation|shared|audit|service|services|offline|endpoint|cache|policy|schema|script|scripts|file|version|worker|variant|variants|offer|reservation|cancellation|review|health|coverage|sort|adapter|contract|operation|dependency|model|status|refresh|matrix|workspace|revalidate|handoff|eligibility|request|response|idempotency|build|log|concurrency|layout|page|package|monolith|microservice|framework|onboarding|state|runbook|infra|docs|stock|code|path|constraint|nullable|fee|evaluator|quantity|item|render|header|body|rule|actor|discriminator|repository|export|generate|filter|desktop|mobile|runtime|upload|parameter|controller|resolve|resolver|hook|balance|rotation|reuse|public|frame|Shop|Provider|Resale|Service|Product|Home|Search|Detail|Success|Orders|Cart|Checkout|Shops|Listings|Needs|executable|visual|taxonomy|form|command|route|Dashboard|metadata|blob|complete|compile|baseline|typecheck|lint|launcher|seed|regression|HEAD|transaction|Staff|retry|revoke|Attribution|app|focus|performance|pilot|workflow|race|geo|membership|server-owned|restore|backup|actionable|confirm|callback|email|cookie|shell|manifest|mode|setting|sync-required|dump|text|uuid|bigint|varchar|timestamptz|boolean|jsonb|geography|geometry|current|snapshot|fixture|credential|ignored|timestamp|full stack|seed fixture|private key|presigned URL|proxy|client local|OAuth credential|Point|B-tree|JSON|JSONB|Git|GitHub Actions|CI/CD|CI|CD|S3|MINIO|ST_DWithin|Transactional|Application Programming Interface|Progressive Web App|Representational State Transfer|Structured Query Language|Entity Relationship Diagram|JavaScript Object Notation|JSON Web Token|Cross-Site Request Forgery|Server-Sent Events|Create Read Update Delete|Application Programming Interface|Hypertext Transfer Protocol|User Interface|User Experience|Continuous Integration|Continuous Delivery'''.split('|')
TERMS += 'char|smallint|rating|comment|Server|CONFIRMED|COMPLETED|test|Method|NEGOTIATED|rollback|README|fulfillment|FREE|FIXED|name|price|CONDITIONAL|demo|me|application|persistence|VND|subtotal|generated|TRACKED|ANY|ALL|issue|order|evaluate|distance|Need|actionability|migrations|routing|Components|Compose|slide|env|tags|E2E|P95|Timeout|AGENTS|API contract|clean verify|Application|Persistence'.split('|')
TERMS = sorted(set(TERMS), key=len, reverse=True)
RX = re.compile(r'(?<![\w])(?:(?i:' + '|'.join(re.escape(x) for x in TERMS) + r')|[A-Za-z][A-Za-z0-9]*_[A-Za-z0-9_]+|[A-Z][a-z]+(?:[A-Z][A-Za-z0-9]*)+|[a-z]+(?:[A-Z][A-Za-z0-9]*)+)(?![\w])')
# Preserve literal paths, URLs, file names, existing quotes, Markdown code and image directives.
PROTECTED = re.compile(r'\bE-(?:HIST|STATIC|VISUAL|PLAN)\b|“[^”]*”|"[^"\n]*"|`[^`]*`|\[\[FIGURE:[^]]+\]\]|https?://[^\s)]+|(?:(?<!\w)/(?!“)[A-Za-z]|\.local/|apps/|packages/|docs/|infra/|scripts/|/api/|/me\b)[A-Za-z0-9_{}./:-]*|\b(?!(?:Next|Node)\.js\b)[A-Za-z0-9_.-]+\.(?:md|java|json|xml|sql|js|ts|png|sh|ps1)\b|\bpnpm\s+(?:dev(?::stop)?|ui:verify-local|install|api-client:check|web:[a-z]+)(?![\w:])|\bapi-client:check\b')

def matches(text):
    protected = [m.span() for m in PROTECTED.finditer(text)]
    return [m for m in RX.finditer(text) if m.group() != 'CoVie' and not any(m.start() < b and m.end() > a for a,b in protected)]

def quote(text):
    for m in reversed(matches(text)):
        text = text[:m.start()] + '“' + m.group() + '”' + text[m.end():]
    return text

def apply_terminology(doc):
    count = 0
    english = False
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1': english = p.text == 'ABSTRACT'
        if english or 'Code' in p.style.name or p.style.name in ('Cover Meta', 'Title'): continue
        count += edit_paragraph(p)
    for ti,table in enumerate(doc.tables):
        for ri,row in enumerate(table.rows):
            for ci,cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    if ti == 0 and ri > 0 and ci == 1 and ' — ' in p.text:
                        nodes=p._p.xpath('.//w:t')
                        value=p.text
                        if nodes and not value.startswith('“'):
                            left,right=value.split(' — ',1)
                            nodes[0].text='“'+left+'” — '+quote(right)
                            for node in nodes[1:]:node.text=''
                    else: count += edit_paragraph(p)
    return count

def edit_paragraph(p):
    nodes=p._p.xpath('.//w:t[not(ancestor::w:fldSimple)]'); original=[t.text or '' for t in nodes]
    text=''.join(original); found=matches(text)
    # Insert at character offsets without flattening runs, bookmarks or page fields.
    starts={m.start() for m in found}; ends={m.end() for m in found}; offset=0
    for node,part in zip(nodes,original):
        out=[]
        for i,c in enumerate(part):
            pos=offset+i
            if pos in starts: out.append('“')
            out.append(c)
            if pos+1 in ends: out.append('”')
        node.text=''.join(out);offset+=len(part)
    return len(found)

if __name__ == '__main__':
    target=BASE/'CoVie_BaoCaoDoAn.docx'
    backup=BASE.parents[1]/'.local/report-before-terminology.docx'
    if not backup.exists(): shutil.copy2(target,backup)
    doc=Document(target); count=apply_terminology(doc);doc.save(target)
    for path in list((BASE/'chapters').glob('*.md'))+[BASE/'front-matter.md']:
        lines=[];fenced=False
        for line in path.read_text(encoding='utf-8').splitlines():
            if line.startswith('```'):fenced=not fenced;lines.append(line)
            else:lines.append(line if fenced else quote(line))
        path.write_text('\n'.join(lines)+'\n',encoding='utf-8')
    (BASE/'CoVie_BaoCaoDoAn.md').write_text('# CoVie — Báo cáo đồ án\n\n'+'\n\n'.join(p.read_text(encoding='utf-8') for p in sorted((BASE/'chapters').glob('*.md'))),encoding='utf-8')
    path=BASE/'navigation-pages.json';data=json.loads(path.read_text(encoding='utf-8'))
    for e in data['entries'].values(): e['text']=quote(e['text'])
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    (BASE/'terminology-review.json').write_text(json.dumps({'quoted_occurrences':count,'terms':TERMS,'exceptions':['English abstract is continuous English prose.','Executable command blocks, URLs, paths and file names retain literal syntax.','Existing quoted passages are not nested.','Original screenshot text is retained as source evidence.']},ensure_ascii=False,indent=2),encoding='utf-8')
    print('Quoted occurrences:',count)
