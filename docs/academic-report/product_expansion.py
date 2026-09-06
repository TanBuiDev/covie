"""Expand the product chapter without regenerating historical technical inventories."""
from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from terminology import quote

BASE=Path(__file__).resolve().parent
CONTENT=json.loads((BASE/'product-expansion-content.json').read_text(encoding='utf-8'))
MANIFEST={e['key']:e for e in json.loads((BASE/'product-image-manifest.json').read_text(encoding='utf-8'))}

def bookmark(p,name,number):
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(number));b.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(number));p._p.insert(0,b);p._p.append(e)

def navigation(doc):
    for p in list(doc.paragraphs):
        if p.style.name=='TOC Entry':p._p.getparent().remove(p._p)
    # Style spelling may be normalized by python-docx.
    for p in list(doc.paragraphs):
        st=p._p.find('w:pPr/w:pStyle',{'w':qn('w:p').split('}')[0][1:]})
        if st is not None and st.get(qn('w:val'))=='TOCEntry':p._p.getparent().remove(p._p)
    paras=doc.paragraphs
    listhead=next(p for p in paras if p.text=='DANH MỤC HÌNH VÀ BẢNG')
    tocend=next(p for p in paras if p._p.xpath('.//w:fldChar[@w:fldCharType="end"]') and not p._p.xpath('.//w:instrText'))
    entries={};headings=[];caps=[]
    for p in doc.paragraphs:
        ishead=p.style.name.startswith('Heading');iscap=p.style.name=='Caption'
        if not(ishead or iscap):continue
        marks=p._p.findall(qn('w:bookmarkStart'))
        if not marks:continue
        name=marks[0].get(qn('w:name'));entries[name]={'text':p.text,'page':1}
        if ishead and p.style.name in ('Heading 1','Heading 2'):headings.append((p.text,name,p.style.name=='Heading 2'))
        if iscap:caps.append((p.text,name,False))
    def nav(title,name,indent):
        p=OxmlElement('w:p');pr=OxmlElement('w:pPr');p.append(pr)
        st=OxmlElement('w:pStyle');st.set(qn('w:val'),'TOCEntry');pr.append(st)
        ind=OxmlElement('w:ind');ind.set(qn('w:left'),'220' if indent else '0');pr.append(ind)
        tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'right');tab.set(qn('w:leader'),'dot');tab.set(qn('w:pos'),'9071');tabs.append(tab);pr.append(tabs)
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),name);r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=title;r.append(t);link.append(r);p.append(link)
        r=OxmlElement('w:r');r.append(OxmlElement('w:tab'));p.append(r)
        f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),' PAGEREF '+name+' \\h ');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);f.append(r);p.append(f)
        return p
    for t,n,i in headings:tocend._p.addprevious(nav(t,n,i))
    anchor=listhead._p
    for t,n,i in sorted(caps,key=lambda x:0 if x[0].startswith('Hình ') else 1):
        new=nav(t,n,i);anchor.addnext(new);anchor=new
    (BASE/'navigation-pages.json').write_text(json.dumps({'source_pdf':None,'entries':entries},ensure_ascii=False,indent=2),encoding='utf-8')

def apply_product_expansion(doc):
    start=next(p._p for p in doc.paragraphs if p.style.name=='Heading 1' and p.text.startswith('CHƯƠNG 5.'))
    end=next(p._p for p in doc.paragraphs if p.style.name=='Heading 1' and p.text.startswith('CHƯƠNG 6.'))
    node=start;old=[];table=None
    while node is not end:
        old.append(node)
        if node.tag==qn('w:tbl') and table is None:table=deepcopy(node)
        node=node.getnext()
    if table is None:
        # A same-sized template is required only for a full regeneration.
        table=deepcopy(next(t._tbl for t in doc.tables if len(t.columns)==3 and len(t.rows)>=7))
    for el in old:el.getparent().remove(el)
    def para(text,style='Normal',break_before=False):
        p=doc.add_paragraph(quote(text),style=style);end.addprevious(p._p)
        p.paragraph_format.page_break_before=break_before
        return p
    ch=para('CHƯƠNG 5. GIỚI THIỆU SẢN PHẨM VÀ GIAO DIỆN', 'Heading 1',True);ch.alignment=1;bookmark(ch,'product_chapter',5000)
    figure_no=0;section_no=0;md=['# '+ch.text]
    used=[]
    for ix,s in enumerate(CONTENT):
        title=s.get('title')
        if title:
            section_no+=1;p=para(title,'Heading 2',ix>0);bookmark(p,'product_section_'+str(section_no),5000+section_no);md+=['','## '+p.text]
        for j,text in enumerate(s.get('before',[])):
            p=para(text,break_before=not title and j==0);md+=['',p.text]
        if s.get('table'):
            p=para('Bảng 5.1. Các không gian chức năng của CoVie','Caption');bookmark(p,'product_table_1',5900)
            end.addprevious(table)
            rows=[['Không gian','Chức năng chính','Kết quả đối với người dùng'],['Người mua','Khám phá, chi tiết và giỏ','Lựa chọn nguồn cung và tạo đơn'],['Người mua','Địa chỉ, đơn hàng và Cần mua','Theo dõi giao dịch và nhu cầu'],['Người bán','Danh mục, điểm bán và chính sách','Tổ chức nguồn cung theo nơi phục vụ'],['Người bán','Đơn hàng, uy tín và thông báo','Theo dõi và xử lý công việc'],['Quản trị','Tài khoản, báo cáo và cấu hình','Kiểm soát hoạt động theo quyền'],['Quản trị','Tra cứu giao dịch và nhật ký','Đối chiếu lịch sử thay đổi']]
            for tr,values in zip(table.findall(qn('w:tr')),rows):
                for tc,value in zip(tr.findall(qn('w:tc')),values):
                    nodes=tc.findall('.//'+qn('w:t'));nodes[0].text=value
                    for t in nodes[1:]:t.text=''
            md+=['','| Không gian | Chức năng chính | Kết quả đối với người dùng |','| --- | --- | --- |']+['| '+' | '.join(r)+' |' for r in rows[1:]]
        if s.get('images'):
            figure_no+=1;keys=s['images'];used+=keys
            p=para('');p.alignment=1;p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
            from PIL import Image
            for j,key in enumerate(keys):
                path=BASE/'assets/product'/(key+'.png');w,h=Image.open(path).size
                width=min(6.0,4.65*w/h) if len(keys)==1 else min(2.55,4.65*w/h)
                if j:p.add_run('   ')
                pic=p.add_run().add_picture(str(path),width=Inches(width));pic._inline.docPr.set('descr',s['caption']+' — '+key)
            cap=f'Hình 5.{figure_no}. '+s['caption'];p=para(cap,'Caption');p.paragraph_format.keep_with_next=True;bookmark(p,'product_figure_'+str(figure_no),6000+figure_no)
            kind=MANIFEST[keys[0]]['kind'];assert all(MANIFEST[k]['kind']==kind for k in keys)
            source='Nguồn: giao diện CoVie với dữ liệu thử nghiệm.' if kind=='application' else 'Nguồn: thiết kế giao diện CoVie trên Figma. [N9]'
            p=para(source,'Source Note');p.paragraph_format.keep_with_next=False
            md+=['',f'[[PRODUCT_FIGURE:{ix}]]','',quote(cap),'',quote(source)]
        for text in s.get('after',[]):p=para(text);md+=['',p.text]
    # Only one overview insertion in chapter 1, even on repeated application.
    p=next(p for p in doc.paragraphs if p.text.startswith('Đề tài “Xây dựng nền tảng'))
    if not any(x.text.startswith('Về trải nghiệm sản phẩm, CoVie') for x in doc.paragraphs):
        new=doc.add_paragraph(quote('Về trải nghiệm sản phẩm, CoVie tổ chức hành trình từ tìm kiếm đến theo dõi giao dịch trong cùng một ứng dụng. Một người mua có thể tìm rau củ, chọn cửa hàng, lưu địa chỉ nhận và tạo đơn; khi chưa có nguồn cung phù hợp, người đó có thể đăng yêu cầu Cần mua. Người bán tiếp nhận đơn theo điểm bán, cập nhật mặt hàng và quản lý điều kiện phục vụ. Quản trị viên có không gian riêng để xem xét tài khoản, phản ánh và lịch sử hoạt động. Phần giới thiệu chi tiết cùng các màn hình của từng vai trò được trình bày tại Chương 5.'));p._p.addnext(new._p)
    # Revise stale references to the former three-screen-only chapter.
    for p in doc.paragraphs:
        old=p.text
        new=old.replace('Các hình trang chủ, chi tiết sản phẩm và tạo đơn sử dụng bộ dữ liệu thử trong môi trường cục bộ. Chúng minh họa giao diện của hành trình người mua, không đại diện cho dữ liệu khách hàng hoặc hoạt động thương mại thực tế. Nguồn hình được ghi dưới từng hình minh họa.', 'Chương 5 sử dụng 29 ảnh giao diện ứng dụng với dữ liệu thử nghiệm và 14 hình thiết kế từ Figma, trình bày trong 31 hình minh họa. Các hình bao quát người mua, người bán, quản trị và những trạng thái giao diện tiêu biểu. Hình thiết kế được phân biệt bằng chú thích nguồn; dữ liệu trong hình không đại diện cho khách hàng hoặc hoạt động thương mại thực tế.').replace('Các hình tại Chương 5 minh họa trang chủ, chi tiết sản phẩm và thông báo tạo đơn.','Các hình tại Chương 5 minh họa các hành trình người mua, người bán, quản trị và những trạng thái giao diện tiêu biểu.')
        if new!=old:
            ts=p._p.xpath('.//w:t');ts[0].text=quote(new)
            for t in ts[1:]:t.text=''
    # All chapter titles keep the user-requested alignment.
    for p in doc.paragraphs:
        if p.style.name=='Heading 1' and p.text.startswith('CHƯƠNG '):p.alignment=1
    navigation(doc)
    (BASE/'chapters/05-trien-khai.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    summary={'sections':section_no,'figure_groups':figure_no,'screenshots':len(used),'application_screens':sum(MANIFEST[k]['kind']=='application' for k in used),'figma_screens':sum(MANIFEST[k]['kind']=='figma' for k in used),'keys':used}
    (BASE/'product-expansion-summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    return summary

if __name__=='__main__':
    path=BASE/'CoVie_BaoCaoDoAn.docx';doc=Document(path);print(apply_product_expansion(doc));doc.save(path)
    body='\n\n'.join(p.read_text(encoding='utf-8') for p in sorted((BASE/'chapters').glob('*.md')))
    (BASE/'CoVie_BaoCaoDoAn.md').write_text('# CoVie — Báo cáo đồ án\n\n'+body,encoding='utf-8')
