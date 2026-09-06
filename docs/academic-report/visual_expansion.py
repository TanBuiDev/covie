"""Add explanatory figures to text-heavy report sections without changing claims."""
from pathlib import Path
import hashlib
import json
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt

from product_expansion import bookmark, navigation

BASE = Path(__file__).resolve().parent
ASSETS = BASE / "assets" / "explanatory"
FONT = "C:/Windows/Fonts/times.ttf"
BOLD = "C:/Windows/Fonts/timesbd.ttf"

FIGURES = [
    {
        "key": "value_journey", "before": "1.3. Đối tượng và phạm vi",
        "caption": "Hình 1.1. Hành trình giá trị chính của nền tảng CoVie",
        "source": "Nguồn: tổng hợp từ phạm vi sản phẩm và hành trình người dùng CoVie. [N1, N9]",
        "kind": "flow", "title": "HÀNH TRÌNH GIÁ TRỊ CHÍNH CỦA COVIE",
        "items": ["Khám phá\nnguồn cung", "Chọn mặt hàng\nhoặc dịch vụ", "Giỏ hàng hoặc\nyêu cầu Cần mua", "Xem trước giá\nvà điều kiện", "Tạo và theo dõi\nđơn hàng", "Đánh giá và\nnhận thông báo"],
        "footer": "Người mua tạo nhu cầu  •  Người bán tổ chức nguồn cung  •  Quản trị viên kiểm soát hoạt động",
        "chapter": "01-tong-quan.md", "marker": "## 1.3. Đối tượng và phạm vi"
    },
    {
        "key": "event_flow", "before": "2.8. “PWA” và giới hạn hoạt động ngoại tuyến",
        "caption": "Hình 2.1. Luồng sự kiện từ giao dịch đến cập nhật giao diện",
        "source": "Nguồn: mô hình hóa từ cơ chế “outbox”, thông báo và “SSE” của CoVie. [N3]",
        "kind": "flow", "title": "LUỒNG SỰ KIỆN VÀ CẬP NHẬT GIAO DIỆN",
        "items": ["Giao dịch nghiệp vụ", "Ghi dữ liệu và\n“outbox”", "Tiến trình đọc\nsự kiện", "Tạo thông báo", "Phát tín hiệu\nqua “SSE”", "Giao diện đọc lại\ndữ liệu mới"],
        "footer": "Dữ liệu nghiệp vụ được xác nhận trong giao dịch; tín hiệu chỉ yêu cầu máy khách đồng bộ lại",
        "chapter": "02-co-so-ky-thuat.md", "marker": "## 2.8. “PWA” và giới hạn hoạt động ngoại tuyến"
    },
    {
        "key": "technology_roles", "before": "CHƯƠNG 3. PHÂN TÍCH YÊU CẦU",
        "caption": "Hình 2.2. Vai trò của các lớp công nghệ trong CoVie",
        "source": "Nguồn: tổng hợp từ cấu trúc “workspace” và cấu hình kỹ thuật của CoVie. [N3, N5]",
        "kind": "layers", "title": "VAI TRÒ CÁC LỚP CÔNG NGHỆ",
        "items": [
            ("TRẢI NGHIỆM", "“Next.js” • “React” • “TypeScript”", "Điều hướng, giao diện, trạng thái sử dụng"),
            ("HỢP ĐỒNG", "“OpenAPI” • “client” sinh tự động", "Đồng bộ kiểu dữ liệu và bề mặt “HTTP”"),
            ("NGHIỆP VỤ", "“Spring Boot” • miền ứng dụng", "Quyền, quy tắc, giao dịch và vòng đời đơn"),
            ("DỮ LIỆU", "“PostgreSQL” • “PostGIS” • “Redis” • “MinIO”", "Lưu giao dịch, không gian, bộ nhớ đệm và ảnh"),
            ("VẬN HÀNH", "“Docker” • “CI” • “runbook”", "Dựng môi trường, kiểm tra và phục hồi")],
        "footer": "Mỗi lớp có trách nhiệm riêng; hợp đồng dữ liệu nối giao diện với nghiệp vụ",
        "chapter": "02-co-so-ky-thuat.md", "marker": "# CHƯƠNG 3. PHÂN TÍCH YÊU CẦU"
    },
    {
        "key": "request_protection", "before": "4.10. Thiết kế giao diện và trạng thái sử dụng",
        "caption": "Hình 4.5. Các lớp bảo vệ một yêu cầu thay đổi dữ liệu",
        "source": "Nguồn: mô hình hóa từ thiết kế xác thực, phân quyền và giao dịch của CoVie. [N4, N7, N11]",
        "kind": "flow", "title": "CÁC LỚP BẢO VỆ YÊU CẦU THAY ĐỔI DỮ LIỆU",
        "items": ["Phiên và\n“CSRF”", "Xác thực\ndanh tính", "Phân quyền theo\nvai trò và điểm bán", "Kiểm tra dữ liệu\nvà trạng thái", "Giao dịch và\n“idempotency”", "Nhật ký và\n“outbox”"],
        "footer": "Máy chủ quyết định quyền và hành động hợp lệ; giao diện không phải nguồn thẩm quyền",
        "chapter": "04-thiet-ke-he-thong.md", "marker": "## 4.10. Thiết kế giao diện và trạng thái sử dụng"
    },
    {
        "key": "test_strategy", "before": "6.2. Phân loại bằng chứng được sử dụng",
        "caption": "Hình 6.1. Chiến lược kiểm thử nhiều lớp của CoVie",
        "source": "Nguồn: tổng hợp từ mã kiểm thử và tài liệu vận hành của CoVie. [N12, N13]",
        "kind": "pyramid", "title": "CHIẾN LƯỢC KIỂM THỬ NHIỀU LỚP",
        "items": [("Diễn tập “staging”", "Triển khai, sao lưu, phục hồi, quay lui"), ("Hành trình đầu cuối", "Luồng người mua và kiểm tra giao diện"), ("Tích hợp", "Cơ sở dữ liệu, giao dịch, hợp đồng “API”"), ("Đơn vị và tĩnh", "Quy tắc miền, kiểu dữ liệu, dựng mã nguồn")],
        "footer": "Kết quả chỉ có giá trị trong phiên bản, dữ liệu và môi trường được ghi nhận",
        "chapter": "06-kiem-thu-danh-gia.md", "marker": "## 6.2. Phân loại bằng chứng được sử dụng"
    },
    {
        "key": "evidence_levels", "before": "6.3. Kết quả lịch sử đã được ghi nhận",
        "caption": "Hình 6.2. Phân biệt bốn nhóm bằng chứng trong báo cáo",
        "source": "Nguồn: quy ước đánh giá được áp dụng trong báo cáo CoVie.",
        "kind": "columns", "title": "BỐN NHÓM BẰNG CHỨNG TRONG BÁO CÁO",
        "items": [("KẾT QUẢ ĐÃ GHI NHẬN", "Có lệnh, thời điểm và kết quả\nDùng để báo cáo trạng thái tại mốc"), ("PHÂN TÍCH TĨNH", "Đọc mã, hợp đồng và cấu hình\nDùng để xác định cơ chế hoặc giới hạn"), ("ẢNH MINH HỌA", "Giao diện với dữ liệu thử nghiệm\nDùng để mô tả trạng thái hiển thị"), ("KIỂM TRA ĐỀ XUẤT", "Kịch bản chưa thực thi\nDùng làm điều kiện nghiệm thu tiếp theo")],
        "footer": "Không dùng ảnh thiết kế hoặc kế hoạch kiểm thử để khẳng định hệ thống đã nghiệm thu",
        "chapter": "06-kiem-thu-danh-gia.md", "marker": "## 6.3. Kết quả lịch sử đã được ghi nhận"
    },
    {
        "key": "roadmap", "before": "7.4. Kết luận",
        "caption": "Hình 7.1. Lộ trình hoàn thiện đề xuất cho CoVie",
        "source": "Nguồn: tổng hợp từ các giới hạn và điều kiện nghiệm thu được xác định trong báo cáo.",
        "kind": "roadmap", "title": "LỘ TRÌNH HOÀN THIỆN ĐỀ XUẤT",
        "items": [("01", "TÍNH ĐÚNG ĐẮN", "Khoảng cách, phạm vi phục vụ, “snapshot”, kiểm thử đồng thời"), ("02", "HỢP ĐỒNG CHỨC NĂNG", "Bổ sung nguồn dữ liệu còn thiếu cho người mua, người bán và quản trị"), ("03", "HỒI QUY VÀ “STAGING”", "Kiểm thử cùng phiên bản, phục hồi, quay lui và số đo tải"), ("04", "THỬ NGHIỆM NGƯỜI DÙNG", "Đo hoàn thành tác vụ và quyết định hướng mở rộng")],
        "footer": "Hoàn thành điều kiện của giai đoạn trước khi mở rộng phạm vi tiếp theo",
        "chapter": "07-ket-luan-phu-luc.md", "marker": "## 7.4. Kết luận"
    },
]


def _fonts():
    return (ImageFont.truetype(BOLD, 48), ImageFont.truetype(BOLD, 30),
            ImageFont.truetype(FONT, 26), ImageFont.truetype(BOLD, 24),
            ImageFont.truetype(FONT, 22))


def _center(draw, box, text, font, fill="#17202A", spacing=8):
    draw.multiline_text(((box[0]+box[2])/2, (box[1]+box[3])/2), text, font=font,
                        fill=fill, anchor="mm", align="center", spacing=spacing)


def _arrow(draw, start, end, color="#0B7A75", width=8):
    draw.line([start, end], fill=color, width=width)
    x, y = end
    draw.polygon([(x, y), (x-22, y-14), (x-22, y+14)], fill=color)


def draw_figure(spec):
    img = Image.new("RGB", (1800, 1000), "white")
    d = ImageDraw.Draw(img)
    title, sub, body, small_bold, small = _fonts()
    d.rectangle((0, 0, 1800, 116), fill="#F1F7F6")
    d.rectangle((0, 0, 18, 116), fill="#0B7A75")
    d.text((900, 60), spec["title"], font=title, fill="#17202A", anchor="mm")
    kind = spec["kind"]
    if kind == "flow":
        items = spec["items"]; n = len(items); gap = 34; margin = 60
        width = (1680-gap*(n-1))/n; y1, y2 = 315, 660
        for i, text in enumerate(items):
            x1 = margin+i*(width+gap); x2 = x1+width
            if i < n-1: _arrow(d, (x2+4, 487), (x2+gap-5, 487))
            d.rounded_rectangle((x1,y1,x2,y2), radius=28, fill="#F8FBFB", outline="#287D78", width=4)
            d.ellipse((x1+width/2-34,y1+34,x1+width/2+34,y1+102), fill="#0B7A75")
            d.text((x1+width/2,y1+68), str(i+1), font=small_bold, fill="white", anchor="mm")
            _center(d,(x1+14,y1+120,x2-14,y2-18),text,sub)
    elif kind == "layers":
        y=170
        colors=["#DDEFEA","#E8F1F7","#F2F2F2","#E8F1F7","#DDEFEA"]
        for i,(label,tech,role) in enumerate(spec["items"]):
            y1=y+i*142; y2=y1+116
            d.rounded_rectangle((105,y1,1695,y2),radius=24,fill=colors[i],outline="#637A80",width=3)
            d.rounded_rectangle((125,y1+16,385,y2-16),radius=16,fill="#164F5B")
            _center(d,(125,y1+16,385,y2-16),label,small_bold,"white")
            d.text((430,(y1+y2)/2-18),tech,font=sub,fill="#17202A",anchor="lm")
            d.text((430,(y1+y2)/2+28),role,font=small,fill="#43545A",anchor="lm")
    elif kind == "pyramid":
        widths=[760,1040,1320,1600]; y=170
        colors=["#CDE9E3","#DCECF2","#E9EFF2","#F2F2F2"]
        for i,(label,desc) in enumerate(spec["items"]):
            w=widths[i]; x1=(1800-w)/2; y1=y+i*168; y2=y1+136
            d.polygon([(x1+48,y1),(1800-x1-48,y1),(1800-x1,y2),(x1,y2)],fill=colors[i],outline="#54727A")
            d.text((900,y1+43),label,font=sub,fill="#164F5B",anchor="mm")
            d.text((900,y1+93),desc,font=small,fill="#33474D",anchor="mm")
    elif kind == "columns":
        colors=["#DDEFEA","#E8F1F7","#F2F2F2","#EAF2EE"]
        n=len(spec["items"]); gap=34; card=(1640-gap*(n-1))/n
        tiny=ImageFont.truetype(FONT,18)
        for i,(label,desc) in enumerate(spec["items"]):
            x1=80+i*(card+gap); x2=x1+card
            d.rounded_rectangle((x1,220,x2,720),radius=30,fill=colors[i],outline="#54727A",width=4)
            cx=(x1+x2)/2; d.ellipse((cx-50,265,cx+50,365),fill="#0B7A75")
            d.text((cx,315),str(i+1),font=sub,fill="white",anchor="mm")
            _center(d,(x1+25,385,x2-25,470),label,sub)
            wrapped='\n'.join(line for part in desc.split('\n') for line in textwrap.wrap(part,26))
            _center(d,(x1+24,475,x2-24,685),wrapped,tiny,"#33474D",5)
    elif kind == "roadmap":
        items=spec["items"]
        d.line((170,470,1630,470),fill="#0B7A75",width=10)
        for i,(num,label,desc) in enumerate(items):
            x=190+i*470
            d.ellipse((x-58,412,x+58,528),fill="#0B7A75")
            d.text((x,470),num,font=sub,fill="white",anchor="mm")
            _center(d,(x-175,185,x+175,385),label,sub)
            _center(d,(x-195,555,x+195,755),"\n".join(textwrap.wrap(desc,25)),small,"#33474D")
            if i<len(items)-1:_arrow(d,(x+70,470),(x+390,470),width=6)
    d.rounded_rectangle((85,835,1715,930),radius=20,fill="#F7F7F7",outline="#D5DDDF",width=2)
    _center(d,(110,845,1690,920),spec["footer"],small,"#43545A")
    ASSETS.mkdir(parents=True,exist_ok=True)
    path=ASSETS/(spec["key"]+".png"); img.save(path,quality=95)
    return path


def _insert_before(doc, anchor, text="", style="Normal"):
    p=doc.add_paragraph(text,style=style);anchor._p.addprevious(p._p);return p


def apply_visual_expansion(doc):
    if any(p.text == FIGURES[0]["caption"] for p in doc.paragraphs):
        return
    manifest=[]
    for number,spec in enumerate(FIGURES,1):
        anchor=next(p for p in doc.paragraphs if p.text==spec["before"])
        path=draw_figure(spec)
        p=_insert_before(doc,anchor);p.alignment=1;p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1
        pic=p.add_run().add_picture(str(path),width=Inches(6.05));pic._inline.docPr.set("descr",spec["caption"])
        cap=_insert_before(doc,anchor,spec["caption"],"Caption");cap.paragraph_format.keep_with_next=True
        bookmark(cap,"explanatory_figure_"+str(number),8000+number)
        src=_insert_before(doc,anchor,spec["source"],"Source Note");src.paragraph_format.keep_with_next=False
        manifest.append({"key":spec["key"],"caption":spec["caption"],"source":spec["source"],
                         "asset":str(path.relative_to(BASE)),"sha256":hashlib.sha256(path.read_bytes()).hexdigest()})
    navigation(doc)
    for spec in reversed(FIGURES):
        chapter=BASE/"chapters"/spec["chapter"]; text=chapter.read_text(encoding="utf-8")
        token=f'[[EXPLANATORY_FIGURE:{spec["key"]}]]'
        if token not in text:
            block=f'{token}\n\n{spec["caption"]}\n\n{spec["source"]}\n\n'
            text=text.replace(spec["marker"],block+spec["marker"],1)
            chapter.write_text(text,encoding="utf-8")
    body="\n\n".join(p.read_text(encoding="utf-8") for p in sorted((BASE/"chapters").glob("*.md")))
    (BASE/"CoVie_BaoCaoDoAn.md").write_text("# CoVie — Báo cáo đồ án\n\n"+body,encoding="utf-8")
    (BASE/"explanatory-figures.json").write_text(json.dumps({"figures":manifest},ensure_ascii=False,indent=2),encoding="utf-8")


if __name__ == "__main__":
    path=BASE/"CoVie_BaoCaoDoAn.docx"; doc=Document(path);apply_visual_expansion(doc);doc.save(path)
    print("Added seven explanatory figures")
