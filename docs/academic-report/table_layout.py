"""Consistent restrained, banded tables for the academic report."""
from pathlib import Path
from docx import Document
from docx.shared import RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def apply_tables(doc):
    for t in doc.tables:
        heads=[c.text for c in t.rows[0].cells]
        if heads[0] in ['Cột','Bước']:
            widths=[2800,2400,3871] if heads[0]=='Cột' else [900,4000,4171]
            for g,w in zip(t._tbl.tblGrid,widths):g.set(qn('w:w'),str(w))
            for row in t.rows:
                for c,w in zip(row.cells,widths):c.width=Twips(w)
        for border in t._tbl.tblPr.findall('.//'+qn('w:tblBorders')+'/*'):
            border.set(qn('w:color'),'D9D9D9');border.set(qn('w:sz'),'4')
        for ri,row in enumerate(t.rows):
            for ci,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=c._tc.get_or_add_tcPr()
                for old in list(pr.findall(qn('w:shd'))):pr.remove(old)
                shd=OxmlElement('w:shd');shd.set(qn('w:val'),'clear')
                shd.set(qn('w:fill'),'414A53' if ri==0 else ('F2F4F6' if ri%2==0 else 'FFFFFF'));pr.append(shd)
                for p in c.paragraphs:
                    p.alignment=A.CENTER if ri==0 or heads[ci] in ['Mã','Bước','Method','Từ viết tắt'] else (A.RIGHT if heads[ci]=='Số CREATE TABLE' else A.LEFT)
                    for r in p.runs:
                        r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
                        if ri==0:r.bold=True
    for p in doc.paragraphs:
        if p.style.name=='Caption' and p.text.startswith('Bảng '):
            p.alignment=A.LEFT

if __name__=='__main__':
    p=Path(__file__).resolve().parent/'CoVie_BaoCaoDoAn.docx'
    d=Document(p);apply_tables(d);d.save(p)
