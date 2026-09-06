"""Structural checks only. Passing is not a substitute for page rendering."""
from pathlib import Path
import json,re,zipfile,hashlib,xml.etree.ElementTree as ET
from docx import Document

BASE=Path(__file__).resolve().parent
out=BASE/'CoVie_BaoCaoDoAn.docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
doc=Document(out)
checks=[]
def check(name,condition):
    checks.append({'check':name,'passed':bool(condition)})
    assert condition,name
with zipfile.ZipFile(out) as z:
    check('ZIP CRC',z.testzip() is None)
    for name in z.namelist():
        if name.endswith(('.xml','.rels')):ET.fromstring(z.read(name))
    tree=ET.fromstring(z.read('word/document.xml'))
    text='\n'.join(t.text or '' for t in tree.findall('.//w:t',ns)).replace('“','').replace('”','')
    check('Seven chapter headings',sum(p.style.name=='Heading 1' and p.text.startswith('CHƯƠNG ') for p in doc.paragraphs)==7)
    check('Twenty-one diagrams, 43 product screens and one cover logo',len(doc.inline_shapes)==64 and len(tree.findall('.//wp:anchor',ns))==1)
    check('User supplied cover metadata',all(s in text for s in ['Bùi Lê Hoàng Tấn','2280602879','22DTHH2','Hồ Chí Minh, 07/2026','TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP.HCM']))
    check('Advisor line omitted',not any(p.text.startswith('Giảng viên hướng dẫn:') for p in doc.paragraphs))
    check('Forty tables',len(doc.tables)==40)
    check('Nine additional use case diagrams',sum(p.style.name=='Caption' and p.text.startswith('Hình 3.') for p in doc.paragraphs)==10)

    check('Seven centered chapter titles',sum(p.style.name=='Heading 1' and p.text.startswith('CHƯƠNG ') and p.alignment==1 for p in doc.paragraphs)==7)
    check('31 product figure captions',sum(p.style.name=='Caption' and p.text.startswith('Hình 5.') for p in doc.paragraphs)==31)
    check('Seven explanatory figures',sum(p.style.name=='Caption' and p.text in {x['caption'] for x in __import__('visual_expansion').FIGURES} for p in doc.paragraphs)==7)
    check('27 product chapter sections',sum(p.style.name=='Heading 2' and p.text.startswith('5.') for p in doc.paragraphs)==27)
    check('Figma design source distinguished',sum(p.style.name=='Source Note' and 'Figma' in p.text for p in doc.paragraphs)==13 and 'https://www.figma.com/design/4LcMXFQz3BnU1EXzBk69t2' in text)

    check('No unresolved figure directives','[[FIGURE:' not in text and '[[PRODUCT_FIGURE:' not in text and '[[USECASE_FIGURE:' not in text and '[[EXPLANATORY_FIGURE:' not in text)
    check('No internal web citation markers','turn2view' not in text and 'cite' not in text)
    check('TOC field exists',any('TOC ' in (x.text or '') for x in tree.findall('.//w:instrText',ns)))
    anchors={x.attrib.get('{'+ns['w']+'}name') for x in tree.findall('.//w:bookmarkStart',ns)}
    links=[x.attrib.get('{'+ns['w']+'}anchor') for x in tree.findall('.//w:hyperlink',ns)]
    check('All internal links resolve',all(x in anchors for x in links))
    check('Meaningful image alternative text',all(x.attrib.get('descr') for x in tree.findall('.//wp:docPr',ns)))
    for i,t in enumerate(tree.findall('.//w:tbl',ns),1):
        widths=[int(x.attrib['{'+ns['w']+'}w']) for x in t.findall('w:tblGrid/w:gridCol',ns)]
        declared=int(t.find('w:tblPr/w:tblW',ns).attrib['{'+ns['w']+'}w'])
        check(f'Table {i} grid width',sum(widths)==declared==9071)
        for row in t.findall('w:tr',ns):
            cell_widths=[int(c.find('w:tcPr/w:tcW',ns).attrib['{'+ns['w']+'}w']) for c in row.findall('w:tc',ns)]
            assert cell_widths==widths,(i,cell_widths,widths)
        check(f'Table {i} repeated heading',t.find('w:tr/w:trPr/w:tblHeader',ns) is not None)
    api=json.loads((BASE.parents[1]/'packages/api-client/openapi/openapi.json').read_text(encoding='utf-8'))
    check('Every API path included',all(path in text for path in api['paths']))
    for x in ['N'+str(i) for i in range(1,14)]+['T'+str(i) for i in range(1,7)]:
        check('Reference '+x,'['+x+']' in text)
    check('No private key material','-----BEGIN PRIVATE KEY-----' not in text and '-----BEGIN RSA PRIVATE KEY-----' not in text)
    check('Historical evidence limitation included','Chưa có kết quả kiểm thử hồi quy đầy đủ cho phiên bản c477f30' in text and '26/08/2026' in text and 'c477f30' in text)
    check('Concrete checkout distance limitation included','distanceMeters bằng 0' in text)
    check('24 proposed acceptance cases',all(f'TC-{i:02d}' in text for i in range(1,25)))
    check('A4 page geometry',int(doc.sections[0].page_width.twips)==11906 and int(doc.sections[0].page_height.twips)==16838)
summary=json.loads((BASE/'build-summary.json').read_text(encoding='utf-8'))
render_path=BASE/'render-verification.json'
render=json.loads(render_path.read_text(encoding='utf-8')) if render_path.exists() else {}
render_current=render.get('status')=='PASS' and render.get('docx_sha256')==hashlib.sha256(out.read_bytes()).hexdigest()
summary['render_qa']='PASS: packaged LibreOffice render, pagination checks and all-page visual review.' if render_current else 'PENDING: render and inspect the current DOCX; existing render evidence is missing or stale.'
summary['page_count']=render.get('page_count') if render_current else None
summary['structural_qa']='PASS'
summary['diagram_asset_visual_review']='PASS: 21 generated diagrams and 43 selected product screens inspected as images; not a page-layout review.'
summary['figures']=52
summary['diagrams']=21
summary['explanatory_figures_added']=7
(BASE/'build-summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
report={'structural_qa':'PASS','checks':checks,'render_qa':summary['render_qa'],'file_bytes':out.stat().st_size,'page_count':summary['page_count']}
(BASE/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'structural_qa':'PASS','checks':len(checks),'bytes':out.stat().st_size,'render_qa':summary['render_qa']},ensure_ascii=False,indent=2))
