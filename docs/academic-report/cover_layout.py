"""Build the CoVie cover from the user-supplied HUTECH cover composition."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent


def apply_cover(doc):
    front = next(p for p in doc.paragraphs if p.text == "LỜI CAM ĐOAN VÀ LỜI CẢM ƠN")
    for item in list(doc.paragraphs):
        if item._p is front._p:
            break
        item._p.getparent().remove(item._p)

    def paragraph(*, before=0, after=0, left=0, align=WD_ALIGN_PARAGRAPH.CENTER,
                  keep=True, style="Cover Meta"):
        p = front.insert_paragraph_before(style=style)
        p.alignment = align
        fmt = p.paragraph_format
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = 1.15
        fmt.left_indent = Cm(left)
        fmt.right_indent = Cm(0)
        fmt.first_line_indent = Cm(0)
        fmt.keep_with_next = keep
        fmt.page_break_before = False
        return p

    def run(p, text, size=13, bold=False):
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.bold = bold
        r.font.color.rgb = RGBColor(0, 0, 0)
        return r

    p = paragraph(after=3, left=4.2)
    run(p, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", 11)
    shape = p.add_run().add_picture(str(BASE / "assets/hutech-logo.png"), width=Cm(3.25))
    shape._inline.docPr.set("descr", "Logo HUTECH chính thức trên bìa báo cáo")
    inline = shape._inline
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" '
        'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>1080000</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>720000</wp:posOffset></wp:positionV>'
        '</wp:anchor>'
    )
    for tag in ("extent", "effectExtent"):
        element = inline.find(qn("wp:" + tag))
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement("wp:wrapNone"))
    for element in list(inline):
        anchor.append(element)
    inline.getparent().replace(inline, anchor)

    p = paragraph(after=0, left=4.2)
    run(p, "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP.HCM", 11, True)

    p = paragraph(before=137, after=89, style="Title")
    run(p, "BÁO CÁO ĐỒ ÁN", 18, True)

    p = paragraph(after=97)
    run(p, "XÂY DỰNG NỀN TẢNG MUA BÁN ĐỊA PHƯƠNG\nCOVIE", 15, True)

    def label_value(label, value, *, left, after=8, size=12, value_bold=True):
        p = paragraph(left=left, after=after, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(6.25))
        run(p, label + ":\t", size)
        run(p, value, size, value_bold)
        return p

    label_value("Ngành", "CÔNG NGHỆ THÔNG TIN", left=1.45)
    label_value("Chuyên ngành", "CÔNG NGHỆ PHẦN MỀM", left=1.45, after=84)
    label_value("Sinh viên thực hiện", "Bùi Lê Hoàng Tấn", left=1.45, value_bold=False)
    label_value("Mã số sinh viên", "2280602879", left=1.45, value_bold=False)
    label_value("Lớp", "22DTHH2", left=1.45, after=0, value_bold=False)

    p = paragraph(before=92, keep=False)
    run(p, "TP. Hồ Chí Minh, 07/2026", 12)

    sect = doc.sections[0]._sectPr
    old = sect.find(qn("w:pgBorders"))
    if old is not None:
        sect.remove(old)
    p = doc.paragraphs[0]
    frames = [(.63, .56, 19.74, 28.58, "1pt"), (.70, .63, 19.60, 28.44, ".5pt")]
    for number, (x, y, width, height, stroke) in enumerate(frames):
        pict = parse_xml(
            f'<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'xmlns:v="urn:schemas-microsoft-com:vml"><v:rect id="cover_frame_{number}" '
            f'style="position:absolute;margin-left:{x}cm;margin-top:{y}cm;width:{width}cm;'
            f'height:{height}cm;z-index:-251654144;mso-position-horizontal-relative:page;'
            f'mso-position-vertical-relative:page" filled="f" strokecolor="#000000" '
            f'strokeweight="{stroke}"/></w:pict>'
        )
        p.add_run()._r.append(pict)


if __name__ == "__main__":
    path = BASE / "CoVie_BaoCaoDoAn.docx"
    document = Document(path)
    apply_cover(document)
    document.save(path)
